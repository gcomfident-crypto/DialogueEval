from __future__ import annotations

import csv
from dataclasses import dataclass
from pathlib import Path


EXCEL_SUFFIXES = {".xlsx", ".xlsm", ".xltx", ".xltm"}
CSV_SUFFIXES = {".csv"}


@dataclass(frozen=True)
class ExtractedEvalStandard:
    source_excel_path: str
    source_sheet: str
    source_row: int
    source_column: int
    output_path: str
    title: str


def is_excel_path(path: str | Path) -> bool:
    return Path(path).suffix.lower() in EXCEL_SUFFIXES


def is_csv_path(path: str | Path) -> bool:
    return Path(path).suffix.lower() in CSV_SUFFIXES


def is_tabular_path(path: str | Path) -> bool:
    return is_excel_path(path) or is_csv_path(path)


def extract_markdown_from_excel(
    excel_path: str | Path,
    output_dir: str | Path,
    *,
    column: int = 2,
    start_row: int = 2,
    sheet_name: str | None = None,
    filename_prefix: str = "eval_standard",
) -> list[ExtractedEvalStandard]:
    if column < 1:
        raise ValueError("column must be 1-based and greater than 0")
    if start_row < 1:
        raise ValueError("start_row must be 1-based and greater than 0")

    try:
        from openpyxl import load_workbook
    except ModuleNotFoundError as exc:
        raise RuntimeError(
            "openpyxl is required to extract Markdown from Excel. "
            "Install dependencies with `pip install -r requirements.txt`."
        ) from exc

    source_path = Path(excel_path)
    destination = Path(output_dir)
    destination.mkdir(parents=True, exist_ok=True)

    workbook = load_workbook(source_path, read_only=True, data_only=True)
    worksheet = workbook[sheet_name] if sheet_name else workbook.active
    extracted: list[ExtractedEvalStandard] = []

    for row_index in range(start_row, worksheet.max_row + 1):
        value = worksheet.cell(row=row_index, column=column).value
        if value is None:
            continue
        markdown = str(value).strip()
        if not markdown:
            continue

        title = first_markdown_title(markdown) or f"row_{row_index}"
        output_path = destination / f"{filename_prefix}_row_{row_index:03d}.md"
        output_path.write_text(markdown.rstrip() + "\n", encoding="utf-8")
        extracted.append(
            ExtractedEvalStandard(
                source_excel_path=str(source_path),
                source_sheet=worksheet.title,
                source_row=row_index,
                source_column=column,
                output_path=str(output_path),
                title=title,
            )
        )

    workbook.close()
    return extracted


def extract_markdown_from_csv(
    csv_path: str | Path,
    output_dir: str | Path,
    *,
    column: int = 2,
    start_row: int = 2,
    filename_prefix: str = "eval_standard",
) -> list[ExtractedEvalStandard]:
    if column < 1:
        raise ValueError("column must be 1-based and greater than 0")
    if start_row < 1:
        raise ValueError("start_row must be 1-based and greater than 0")

    source_path = Path(csv_path)
    destination = Path(output_dir)
    destination.mkdir(parents=True, exist_ok=True)

    extracted: list[ExtractedEvalStandard] = []
    with source_path.open("r", encoding="utf-8-sig", newline="") as file:
        reader = csv.reader(file)
        for row_index, row in enumerate(reader, start=1):
            if row_index < start_row or len(row) < column:
                continue
            markdown = str(row[column - 1]).strip()
            if not markdown:
                continue

            title = first_markdown_title(markdown) or f"row_{row_index}"
            output_path = destination / f"{filename_prefix}_row_{row_index:03d}.md"
            output_path.write_text(markdown.rstrip() + "\n", encoding="utf-8")
            extracted.append(
                ExtractedEvalStandard(
                    source_excel_path=str(source_path),
                    source_sheet="",
                    source_row=row_index,
                    source_column=column,
                    output_path=str(output_path),
                    title=title,
                )
            )
    return extracted


def first_markdown_title(markdown: str) -> str:
    for line in markdown.splitlines():
        stripped = line.strip()
        if stripped.startswith("#"):
            return stripped.lstrip("#").strip()
    return ""
